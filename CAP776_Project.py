from openpyxl import load_workbook
from datetime import date, timedelta
file_name = "CAP776_Activity_20Aug_to_25Sep_2026.xlsx"

try:
    workbook = load_workbook(file_name)
    sheet = workbook["Activity Data"]

    print("Excel file successfully opened!")
    print("Rows:", sheet.max_row)
    print("Columns:", sheet.max_column)

except FileNotFoundError:
    print("Excel file nahi mili. Excel aur Python file same folder mein rakho.")

except Exception as e:
    print("Error:", e)

dates = []

for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, min_col=1, max_col=1, values_only=True):
    if row[0]:
        dates.append(row[0].date())

start_date = date(2026, 8, 20)
end_date = date(2026, 9, 25)

expected_dates = {
    start_date + timedelta(days=i)
    for i in range((end_date - start_date).days + 1)
}

recorded_dates = set(dates)
missing_dates = sorted(expected_dates - recorded_dates)
duplicate_count = len(dates) - len(recorded_dates)

print("\nDate Validation:")
print("Expected days:", len(expected_dates))
print("Recorded days:", len(recorded_dates))
print("Missing days:", len(missing_dates))
print("Duplicate dates:", duplicate_count)
print("Missing dates:", missing_dates)
print("\nData Validation:")

invalid_rows = []

for row_num in range(6, sheet.max_row + 1):
    values = [sheet.cell(row_num, col).value for col in range(1, 14)]

    date_value = values[0]
    numeric_values = values[1:10]
    feeling = values[10]
    satisfaction = values[11]
    energy = values[12]

    if date_value is None:
        invalid_rows.append((row_num, "Missing date"))
        continue

    if any(v is None or not isinstance(v, (int, float)) for v in numeric_values):
        invalid_rows.append((row_num, "Missing/invalid numeric value"))
        continue

    if any(v < 0 for v in numeric_values):
        invalid_rows.append((row_num, "Negative value"))

    if feeling not in [1, 2, 3, 4, 5]:
        invalid_rows.append((row_num, "Invalid feeling"))

    if satisfaction not in [1, 2, 3, 4, 5]:
        invalid_rows.append((row_num, "Invalid satisfaction"))

    if energy not in [1, 2, 3]:
        invalid_rows.append((row_num, "Invalid energy"))

print("Invalid records:", len(invalid_rows))
print("Invalid details:", invalid_rows)
print("\n--- AVERAGES & INDICES ---")

data = []

for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, values_only=True):
    if row[0] is not None:
        data.append(row)

n = len(data)

avg_sleep = sum(row[1] for row in data) / n
avg_fitness = sum(row[2] for row in data) / n
avg_study = sum(row[3] for row in data) / n
avg_coding = sum(row[4] for row in data) / n
avg_class = sum(row[5] for row in data) / n
avg_other = sum(row[7] for row in data) / n
avg_total = sum(row[8] for row in data) / n
avg_free = sum(row[9] for row in data) / n

avg_feeling = sum(row[10] for row in data) / n
avg_satisfaction = sum(row[11] for row in data) / n
avg_energy = sum(row[12] for row in data) / n

TPI = avg_coding
AAI = avg_study + avg_class
PhAI = avg_fitness
SRI = avg_sleep
ABI = avg_free
TUI = avg_total
EI = (avg_feeling + avg_satisfaction + avg_energy) / 3
DCI = (len(recorded_dates) / len(expected_dates)) * 100

PAI = (
    0.15 * TPI
    + 0.20 * AAI
    + 0.15 * PhAI
    + 0.20 * SRI
    + 0.15 * TUI
    + 0.10 * EI
    + 0.05 * DCI
)

print("Average Sleep:", round(avg_sleep, 2))
print("Average Fitness:", round(avg_fitness, 2))
print("Average Study:", round(avg_study, 2))
print("Average Coding:", round(avg_coding, 2))
print("Average Class:", round(avg_class, 2))
print("Average Other:", round(avg_other, 2))
print("Average Free:", round(avg_free, 2))

print("\nTPI:", round(TPI, 2))
print("AAI:", round(AAI, 2))
print("PhAI:", round(PhAI, 2))
print("SRI:", round(SRI, 2))
print("ABI:", round(ABI, 2))
print("TUI:", round(TUI, 2))
print("EI:", round(EI, 2))
print("DCI:", round(DCI, 2))
print("PAI:", round(PAI, 2))
study_values = []
class_values = []
other_values = []

for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row, values_only=True):
    if row[3] is not None:
        study_values.append(float(row[3]))
    if row[5] is not None:
        class_values.append(float(row[5]))
    if row[7] is not None:
        other_values.append(float(row[7]))

print("Average Study/day:", round(sum(study_values) / len(study_values), 2))
print("Average Class/day:", round(sum(class_values) / len(class_values), 2))
print("Average Other Activities/day:", round(sum(other_values) / len(other_values), 2))